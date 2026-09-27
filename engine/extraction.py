"""Bounded read-only extraction. Formulas are data, never Python/shell code."""
import csv,io,re,uuid,zipfile,os
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
from datetime import date,datetime
from .common import EngineError,number,output
from .aliases import concept,DEFINITIONS,normalize
from . import __version__

EXPENSES={'cogs','gna','sga','depreciation_amortization','finance_cost','capex','selling_expenses','r_and_d','ecl_expense','impairment','customer_service_cost'}
MAX_CELLS=250000

def parse_number(value):
    if value is None or isinstance(value,bool):return None
    if isinstance(value,(int,float,Decimal)):
        try:return number(value)
        except EngineError:return None
    text=str(value).strip().translate(str.maketrans('٠١٢٣٤٥٦٧٨٩٫٬','0123456789.,'))
    if not text or text.lower() in ('n/a','na','null','none','—','–','-'):return None
    negative=text.startswith('(') and text.endswith(')')
    text=text.strip('()').replace(',','').replace(' ','').replace('−','-')
    text=re.sub(r'^(SAR|USD|EUR|ر\.س\.?|﷼|\$)','',text,flags=re.I).rstrip('%٪')
    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)',text):return None
    try:return number(text)*(-1 if negative else 1)
    except EngineError:return None

def year(value):
    if isinstance(value,(datetime,date)):return str(value.year)
    text=str(value or '').strip()
    # Reject paragraphs mentioning years: headers, not narrative inference.
    if len(text)>45:return None
    m=re.fullmatch(r'(?:FY\s*)?((?:19|20)\d{2})(?:\.0)?',text,re.I)
    if m:return m.group(1)
    if re.search(r'december|dec\b|ديسمبر|31[/-]12|12[/-]31',text,re.I):
        m=re.search(r'\b((?:19|20)\d{2})\b',text)
        return m.group(1) if m else None
    return None

class SafeFormula:
    """Whitelisted arithmetic only, bounded refs, no eval or external I/O."""
    def __init__(self,wb):self.wb=wb;self.memo={};self.active=set()
    def get(self,sheet,addr):
        from openpyxl.utils.cell import range_boundaries,get_column_letter
        addr=addr.replace('$','')
        if '[' in addr or ']' in addr:raise ValueError('External reference blocked')
        if '!' in addr:sheet,addr=addr.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
        if sheet not in self.wb.sheetnames:raise ValueError('Unknown worksheet')
        if ':' in addr:
            c1,r1,c2,r2=range_boundaries(addr)
            if (c2-c1+1)*(r2-r1+1)>10000:raise ValueError('Formula range limit')
            return [self.get(sheet,f'{get_column_letter(c)}{r}') for r in range(r1,r2+1) for c in range(c1,c2+1)]
        if not re.fullmatch(r'[A-Z]{1,3}[1-9]\d{0,6}',addr):raise ValueError('Unsupported reference')
        key=(sheet,addr)
        if key in self.memo:return self.memo[key]
        if key in self.active or len(self.active)>100:raise ValueError('Cyclic/deep formula')
        cell=self.wb[sheet][addr]
        if cell.data_type!='f':
            result=parse_number(cell.value)
            if result is None:raise ValueError('Missing/non-numeric precedent')
            return result
        self.active.add(key)
        try:result=self.evaluate(sheet,cell.value);self.memo[key]=result;return result
        finally:self.active.remove(key)
    def evaluate(self,sheet,formula):
        from openpyxl.formula.tokenizer import Tokenizer
        if len(formula)>4096 or '[' in formula or ']' in formula:raise ValueError('Formula blocked')
        tokens=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE'];position=0
        precedence={'+':1,'-':1,'*':2,'/':2}
        def parse(minimum=0):
            nonlocal position
            if position>=len(tokens):raise ValueError('Incomplete formula')
            t=tokens[position];position+=1
            if t.type=='OPERAND' and t.subtype=='NUMBER':left=number(t.value)
            elif t.type=='OPERAND' and t.subtype=='RANGE':left=self.get(sheet,t.value)
            elif t.type=='OPERATOR-PREFIX' and t.value in ('+','-'):left=parse(3)*(-1 if t.value=='-' else 1)
            elif t.type=='PAREN' and t.subtype=='OPEN':
                left=parse()
                if position>=len(tokens) or tokens[position].subtype!='CLOSE':raise ValueError('Missing close')
                position+=1
            elif t.type=='FUNC' and t.subtype=='OPEN':
                name=t.value[:-1].upper()
                if name not in ('SUM','ROUND','ABS','AVERAGE','MIN','MAX'):raise ValueError('Unsupported function')
                args=[]
                while position<len(tokens) and tokens[position].subtype!='CLOSE':
                    a=parse();args.extend(a if isinstance(a,list) else [a])
                    if position<len(tokens) and tokens[position].type=='SEP':position+=1
                    else:break
                if position>=len(tokens) or tokens[position].subtype!='CLOSE':raise ValueError('Missing close')
                position+=1
                if not args:raise ValueError('Missing function inputs')
                if name=='SUM':left=sum(args,Decimal(0))
                elif name=='ROUND':
                    if len(args)!=2 or args[1]!=args[1].to_integral_value() or abs(args[1])>15:raise ValueError('Invalid ROUND')
                    left=args[0].quantize(Decimal(1).scaleb(-int(args[1])),rounding=ROUND_HALF_UP)
                elif name=='ABS':
                    if len(args)!=1:raise ValueError('Invalid ABS')
                    left=abs(args[0])
                elif name=='AVERAGE':left=sum(args)/len(args)
                elif name=='MIN':left=min(args)
                else:left=max(args)
            else:raise ValueError('Unsupported formula token')
            while position<len(tokens):
                t=tokens[position]
                if t.type!='OPERATOR-INFIX' or t.value not in precedence or precedence[t.value]<minimum:break
                position+=1;right=parse(precedence[t.value]+1)
                if t.value=='+':left+=right
                elif t.value=='-':left-=right
                elif t.value=='*':left*=right
                else:left/=right
            return left
        result=parse()
        if position!=len(tokens) or isinstance(result,list):raise ValueError('Unsupported formula remainder')
        return number(result)

def resolve_scale(rows,requested,warnings):
    if requested is not None:
        scale=number(requested,False)
        if scale<=0 or scale>Decimal('1000000000'):raise EngineError('Invalid currency scale')
        return scale
    text=' '.join(str(c or '') for row in rows[:15] for c in row).lower()
    if re.search(r"million|مليون",text):return Decimal(1000000)
    if re.search(r"thousand|'000|’000|الف|ألف",text):return Decimal(1000)
    if re.search(r'all (?:figures|amounts) in (?:sar|saudi riyals)|جميع المبالغ بالريال',text):return Decimal(1)
    warnings.append('الوحدة غير مصرح بها؛ افترضت وحدة واحدة. راجع scale قبل الاعتماد. / Currency scale was not declared; assumed 1 pending review.')
    return Decimal(1)

def make_fact(label,period,value,source,scale,currency,raw_text=None,formula=None,cache_status=None,raw_cached=None,number_format=''):
    key=concept(label) or 'unmapped'
    ar,en,_=DEFINITIONS.get(key,(str(label),str(label),''))
    original=parse_number(value)
    unit='percent' if '%' in number_format or (raw_text and str(raw_text).rstrip().endswith(('%','٪'))) else 'currency'
    if key in ('peer_count','deposits_eligible'):unit='number'
    if key in ('peer_gross_margin',):unit='percent'
    if key in ('covenant_limit','covenant_actual'):unit='number'
    factor=scale if unit=='currency' else Decimal(1)
    normalized=original*factor if original is not None else None
    if unit=='percent' and '%' in number_format and normalized is not None:normalized*=100
    if key in EXPENSES and normalized is not None:normalized=abs(normalized)
    fact={'id':str(uuid.uuid5(uuid.NAMESPACE_URL,f'{source}|{period}|{label}')),'concept':key,'label_ar':ar,'label_en':en,'period':str(period),'value':output(normalized),'currency':currency,'unit':unit,'scale':output(factor),'original_value':output(original),'raw_text':str(raw_text if raw_text is not None else value) if raw_text is not None or value is not None else None,'source':source,'review_status':'needs_review','scope':'consolidated','sign_convention':'expense magnitude; original preserved' if key in EXPENSES else 'reported sign'}
    if formula:fact.update(formula=formula,cache_status=cache_status,cached_value=raw_cached,formula_engine=f'basira-safe-arithmetic/{__version__}')
    return fact

def extract_xlsx(path,currency,requested):
    import openpyxl
    warnings=[]
    with zipfile.ZipFile(path) as archive:
        entries=archive.infolist()
        if len(entries)>10000 or sum(x.file_size for x in entries)>150_000_000:raise EngineError('Expanded workbook too large','file_limit')
        if any(x.file_size>50_000_000 for x in entries):raise EngineError('Workbook entry too large','file_limit')
        if any('vbaProject' in x.filename for x in entries):warnings.append('Macros detected and ignored; never executed.')
        if any('externalLinks/' in x.filename for x in entries):warnings.append('External workbook links blocked; references need review.')
    wb=openpyxl.load_workbook(path,data_only=False,read_only=False,keep_links=False)
    cached=openpyxl.load_workbook(path,data_only=True,read_only=False,keep_links=False)
    total=sum(s.max_row*s.max_column for s in wb)
    if total>MAX_CELLS:raise EngineError('Workbook cell limit exceeded','file_limit')
    calculator=SafeFormula(wb);facts=[];preview=[];source_cells=[]
    for ws in wb:
        rows=list(ws.iter_rows())
        values=[[c.value for c in row] for row in rows]
        score=sum(concept(c.value) is not None for row in rows for c in row if isinstance(c.value,str))
        # Analysis/assumption sheets remain previewable, never merged into source facts.
        excluded=bool(re.search(r'assumption|analysis|model|recommendation|benchmark|افتراض|تحليل',ws.title,re.I))
        scale=resolve_scale(values,requested,warnings) if score and not excluded else Decimal(1)
        year_columns={}
        for row_index,row in enumerate(rows,1):
            if len(preview)<100:preview.append({'sheet':ws.title,'row':row_index,'text':' | '.join(f'{c.coordinate}: {str(c.value)[:300]}' for c in row if c.value is not None)[:1500]})
            for c in row:
                if c.value is not None and len(source_cells)<5000:source_cells.append({'sheet':ws.title,'cell':c.coordinate,'raw':str(c.value) if c.data_type=='f' else c.value,'cached':cached[ws.title][c.coordinate].value if c.data_type=='f' else None,'format':c.number_format})
            headers={i:year(c.value) for i,c in enumerate(row) if year(c.value)}
            if len(headers)>=2 or (headers and any(normalize(c.value) in ('line item','metric','البند','item','concept') for c in row if c.value)):
                year_columns=headers;continue
            if excluded or not year_columns:continue
            start=min(year_columns)
            labels=[str(c.value).strip() for c in row[:start] if isinstance(c.value,str) and not c.value.startswith('=')]
            if not labels:continue
            label=labels[-1]
            if re.match(r'check|section|total liabilities.*equity|balance check',label,re.I):continue
            key=concept(label)
            # Never manufacture facts for an entirely blank heading.
            if not key and not any(parse_number(row[i].value) is not None for i in year_columns if i<len(row)):continue
            for index,period in year_columns.items():
                if index>=len(row):continue
                c=row[index];value=c.value;cache_state=None
                raw_cache=cached[ws.title][c.coordinate].value if c.data_type=='f' else None
                if c.data_type=='f':
                    try:
                        value=calculator.get(ws.title,c.coordinate)
                        cache_number=parse_number(raw_cache)
                        if cache_number is None:cache_state='recalculated_no_cache'
                        elif abs(value-cache_number)<=max(Decimal('.000001'),abs(value)*Decimal('1e-10')):cache_state='independently_verified'
                        else:cache_state='cache_mismatch';warnings.append(f'{ws.title}!{c.coordinate}: saved formula cache differs from deterministic recalculation.')
                    except Exception:
                        value=None;cache_state='unsupported_or_missing_precedent'
                        warnings.append(f'{ws.title}!{c.coordinate}: formula not evaluated; mapped value unavailable. Cached value retained for review only.')
                if c.data_type=='e':value=None;warnings.append(f'{ws.title}!{c.coordinate}: Excel error {c.value}')
                facts.append(make_fact(label,period,value,{'sheet':ws.title,'cell':c.coordinate,'row':row_index,'label_cell':row[start-1].coordinate if start else None},scale,currency,c.value,c.value if c.data_type=='f' else None,cache_state,raw_cache,c.number_format))
    wb.close();cached.close()
    return finish(facts,warnings,preview,'xlsx',source_cells)

def tabular(rows,currency,requested,kind,page=None):
    warnings=[];scale=resolve_scale(rows,requested,warnings);facts=[];headers={};preview=[]
    if rows:
        names={normalize(v):i for i,v in enumerate(rows[0])}
        if {'period','concept','value'}.issubset(names):
            for rownum,row in enumerate(rows[1:],2):
                if len(row)<=max(names.values()):continue
                p=year(row[names['period']]);label=row[names['concept']]
                if p:facts.append(make_fact(label,p,row[names['value']],{'row':rownum},scale,currency))
            return finish(facts,warnings,[{'text':str(r)[:1000]} for r in rows[:80]],kind)
    for rownum,row in enumerate(rows,1):
        if rownum<=80:preview.append({'page':page,'row':rownum,'text':' | '.join(map(str,row))[:1500]})
        hs={i:year(v) for i,v in enumerate(row) if year(v)}
        # A genuine tabular header has year cells PLUS at least one non-year cell (the label
        # column, even if its header text is blank/generic). A row where every single cell is
        # a year is far more likely a disconnected fragment of a multi-line header block — real
        # PDFs routinely print a header like "31 December" / "2024" as separate physical lines
        # that land as their own rows once split by line — and trusting it sets `first=0`,
        # which then slices every following data row's label to an empty string via
        # `row[:first]`, silently discarding real data rows immediately after it (confirmed
        # against a real annual report: this was the exact reason a page with a clean,
        # correctly-columned "Revenue | note | 2024 | 2023" row produced zero facts).
        if (len(hs)>=2 and len(hs)<len(row)) or (hs and any(normalize(v) in ('line item','item','البند','concept') for v in row)):headers=hs;continue
        if not headers:continue
        first=min(headers);label=' '.join(str(v) for v in row[:first]).strip()
        if not label:continue
        for i,p in headers.items():
            if i<len(row):facts.append(make_fact(label,p,row[i],{'page':page,'row':rownum} if page else {'row':rownum},scale,currency))
    return finish(facts,warnings,preview,kind)

def pdf_tabular(rows,currency,requested,page):
    """PDF-specific counterpart to `tabular()`, kept fully separate so CSV/XLSX's rigid,
    well-tested grid parsing is never touched. Real PDFs routinely print a table's header as
    several disconnected physical-line fragments (e.g. "31 December"/"31 December" then
    "Declaration"/"Note" then "2024"/"2023" as three separate rows), so no single row reliably
    carries both the period labels and a matching column count/position for the data rows
    that follow — index-based lookup (`row[i] for i in headers`, `tabular()`'s approach) then
    silently drops or misaligns real data. Confirmed against a real annual report: a page with
    a clean, correctly-columned "Revenue | note | 2024 | 2023" data row produced zero facts
    under index alignment because the header fragment above it had no leading label cell.
    Here, a header row only needs to name the ordered set of periods; each data row's LAST
    len(periods) cells are taken as that row's values regardless of the header's own column
    count, and everything before them is the label. This matches how real statement rows are
    actually laid out (label, optional note reference, then trailing values) far more robustly
    than index alignment.
    """
    warnings=[];scale=resolve_scale(rows,requested,warnings);facts=[];periods=None;preview=[]
    for rownum,row in enumerate(rows,1):
        if rownum<=80:preview.append({'page':page,'row':rownum,'text':' | '.join(map(str,row))[:1500]})
        hs=[year(v) for v in row if year(v)]
        if len(hs)>=2:periods=hs;continue
        if not periods:continue
        n=len(periods)
        if len(row)<=n:continue
        label_cells=list(row[:-n])
        # Real statement rows commonly carry a bare note-reference number (e.g. "25", "26")
        # as its own column between the label and the values ("Revenue | 25 | 2,299,... |
        # 2,133,..."). Left in the label it silently breaks concept matching ("Revenue 25"
        # matches nothing) even though the row is otherwise clean — confirmed directly: this
        # was the one remaining reason "Revenue" itself stayed unmapped after every other
        # income-statement line on the same real page mapped correctly.
        if len(label_cells)>1 and re.fullmatch(r'\d{1,3}',str(label_cells[-1]).strip()):label_cells.pop()
        label=' '.join(str(v) for v in label_cells).strip()
        if not label:continue
        for p,v in zip(periods,row[-n:]):facts.append(make_fact(label,p,v,{'page':page,'row':rownum},scale,currency))
    return finish(facts,warnings,preview,'pdf')

def _cluster_rows(words):
    """Group already-selected words (one physical column/page region) into visual lines by
    y-position, then into cells by x-gap, using the ~3pt/~8pt thresholds established from
    real financial-statement word spacing (see pdf_rows)."""
    words=sorted(words,key=lambda w:(w['top'],w['x0']))
    lines=[]
    for w in words:
        if lines and abs(w['top']-lines[-1][0]['top'])<=3:lines[-1].append(w)
        else:lines.append([w])
    rows=[]
    for line in lines:
        line=sorted(line,key=lambda w:w['x0'])
        cols=[[line[0]]]
        for prev,w in zip(line,line[1:]):
            if w['x0']-prev['x1']>8:cols.append([w])
            else:cols[-1].append(w)
        rows.append([' '.join(c['text'] for c in group) for group in cols])
    return rows

def pdf_rows(page):
    """Position-aware row/column reconstruction for one PDF page, replacing a prior
    plain-text-then-regex approach that broke on real financial statements: two real
    corporate PDFs tested (both born-digital, not scanned) showed intra-label word gaps of
    ~3pt versus real column gaps of >=16pt, so a fixed 8pt threshold cleanly separates a
    multi-word label from adjacent note/period columns without splitting the label itself.
    A second real failure mode: wide (landscape) pages that lay out two separate statements
    side by side print both at similar y-coordinates, so naive top-to-bottom line grouping
    interleaves their rows into one incoherent stream. When the page is wide and word x0
    positions show one dominant gap (a real column gutter, not just inter-word spacing), each
    side is clustered into rows independently and concatenated, left side first.
    Returns (rows, note); note is None, 'image_only' (real signal: page has zero extractable
    characters but does contain an image, i.e. very likely a scanned page — OCR is still not
    implemented, so this must be reported, never silently treated as an empty page) or
    'error'.
    """
    try:
        words=page.extract_words(use_text_flow=False,keep_blank_chars=False)
    except Exception:
        return [],'error'
    if not words:
        return [],('image_only' if getattr(page,'images',None) else None)
    if sum(len(w['text']) for w in words)>1_000_000:raise EngineError('PDF text limit exceeded','file_limit')
    if getattr(page,'width',0)>900:
        xs=sorted(w['x0'] for w in words)
        gaps=[(xs[i+1]-xs[i],xs[i],xs[i+1]) for i in range(len(xs)-1)]
        if gaps:
            gap,left_edge,right_edge=max(gaps)
            if gap>60:
                mid=(left_edge+right_edge)/2
                left=[w for w in words if w['x0']<mid]
                right=[w for w in words if w['x0']>=mid]
                if left and right:
                    return _cluster_rows(left)+_cluster_rows(right),None
    return _cluster_rows(words),None

def finish(facts,warnings,preview,kind,cells=None):
    unmapped=any(f['concept']=='unmapped' for f in facts) or not facts
    if unmapped:warnings.append('بعض البنود تحتاج ربطًا يدويًا. / Some fields require manual mapping.')
    return {'engine_version':__version__,'facts':facts,'periods':sorted({f['period'] for f in facts}),'warnings':list(dict.fromkeys(warnings))[:200],'source_preview':preview,'source_cells':cells or [],'document_type':kind,'requires_manual_mapping':unmapped}

def extract(request):
    path=Path(str(request.get('path','')))
    if not path.is_absolute() or not path.is_file():raise EngineError('Absolute server-owned file path required')
    if path.stat().st_size>10*1024*1024:raise EngineError('File exceeds10MB','file_limit')
    currency=str(request.get('currency') or 'SAR').upper()
    if not re.fullmatch(r'[A-Z]{3}',currency):raise EngineError('Currency must be a three-letter code')
    extension=Path(str(request.get('filename') or path.name)).suffix.lower()
    scale=request.get('scale')
    if extension=='.xlsx':return extract_xlsx(path,currency,scale)
    if extension in ('.csv','.tsv'):
        text=path.read_text(encoding='utf-8-sig')
        try:dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
        except csv.Error:dialect=csv.excel_tab if extension=='.tsv' else csv.excel
        rows=list(csv.reader(io.StringIO(text),dialect))
        if sum(map(len,rows))>MAX_CELLS:raise EngineError('CSV cell limit exceeded','file_limit')
        return tabular(rows,currency,scale,'csv')
    if extension=='.pdf':
        from pypdf import PdfReader
        import pdfplumber
        reader=PdfReader(path)
        if reader.is_encrypted:raise EngineError('Encrypted PDF is not supported','encrypted_file')
        if len(reader.pages)>300:raise EngineError('PDF page limit exceeded','file_limit')
        facts=[];warnings=[];preview=[]
        with pdfplumber.open(path) as pdf:
            for page_number,page in enumerate(pdf.pages,1):
                rows,note=pdf_rows(page)
                if note=='image_only':
                    warnings.append(f'Page {page_number}: no extractable text but contains an image; likely a scanned page. OCR is not enabled — this page requires manual entry, not an automatic zero.')
                elif note=='error':
                    warnings.append(f'Page {page_number}: text extraction unavailable; manual source review required.')
                preview.append({'page':page_number,'text':'\n'.join(' | '.join(r) for r in rows)[:12000]})
                result=pdf_tabular(rows,currency,scale,page_number)
                for fact in result['facts']:fact['source']['bbox_status']='unavailable_text_line_only'
                facts.extend(result['facts']);warnings.extend(result['warnings'])
        warnings.append('PDF text/table extraction is provisional and requires source-page review. OCR and scanned-page recognition are not enabled.')
        warnings.append('PDF references identify page and extracted text row only; bounding-box coordinates are unavailable.')
        return finish(facts,warnings,preview[:100],'pdf')
    raise EngineError('Supported input types are .xlsx, .csv, .tsv and text .pdf','unsupported_file')
