import openpyxl,json,hashlib,zipfile,sys
from pathlib import Path
p=Path(sys.argv[1]) if len(sys.argv)>1 else Path('/Users/asuayri/Downloads/Elm_Raw_Data_FY2023-FY2025 (2).xlsx')
out=Path(__file__).resolve().parent
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
result={'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'properties':str(wf.properties),'calculation':str(wf.calculation),'defined_names':str(wf.defined_names),'sheets':[]}
lines=[]
for s in wf:
    d={'name':s.title,'dimensions':s.calculate_dimension(),'state':s.sheet_state,'merged_ranges':[str(x) for x in s.merged_cells.ranges],'hidden_rows':[k for k,v in s.row_dimensions.items() if v.hidden],'hidden_columns':[k for k,v in s.column_dimensions.items() if v.hidden],'cells':[],'charts':len(s._charts),'images':len(s._images),'tables':list(s.tables),'autofilter':s.auto_filter.ref}
    lines.append('\n### '+s.title+' '+d['dimensions'])
    for row in s:
        r=[]
        for c in row:
            if c.value is not None or c.comment or c.hyperlink:
                item={'cell':c.coordinate,'value':c.value,'type':c.data_type,'cached':wv[s.title][c.coordinate].value,'format':c.number_format,'comment':c.comment.text if c.comment else None,'comment_author':c.comment.author if c.comment else None,'hyperlink':c.hyperlink.target if c.hyperlink else None}
                d['cells'].append(item)
                r.append(f"{c.coordinate}: {c.value}"+(f" [CACHED: {item['cached']}]" if c.data_type=='f' else '')+(f" [COMMENT: {item['comment']}]" if item['comment'] else ''))
        if r: lines.append(' | '.join(r))
    result['sheets'].append(d)
out.joinpath('full_cell_extraction.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str))
out.joinpath('workbook_all_cells.txt').write_text('\n'.join(lines))
print(json.dumps([{k:v for k,v in s.items() if k not in ('cells','autofilter')}|{'populated_cells':len(s['cells']),'formulas':sum(c['type']=='f' for c in s['cells']),'comments':sum(bool(c['comment']) for c in s['cells'])} for s in result['sheets']],ensure_ascii=False,indent=2))
